"""Menu import (P1-4): a workbook in, the brand's menu upserted through the same catalog functions
the back-office uses, so every row is versioned for devices exactly like a hand edit. The caller
commits only when nothing was wrong (one transaction: all rows or none)."""

import io
from collections.abc import Callable, Container
from dataclasses import asdict, dataclass, field
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from poslite.catalog.application.schemas import (
    CategoryIn,
    ChannelPrice,
    ChannelPricesIn,
    ModifierGroupIn,
    OptionIn,
    OutletProductIn,
    ProductIn,
)
from poslite.catalog.domain.menu_import import SHEETS, MenuImport, Row, RowError, parse_menu
from poslite.catalog.infrastructure import channels
from poslite.catalog.infrastructure import service as catalog
from poslite.catalog.infrastructure.models import (
    CategoryModel,
    ModifierGroupModel,
    ModifierOptionModel,
    ProductModel,
)
from poslite.shared.domain.errors import DomainError
from poslite.shared.infrastructure.unit_of_work import SqlAlchemyUnitOfWork

COLUMNS = {
    "categories": ["name", "print_role", "sort_order"],
    "modifier_groups": ["group", "option", "price", "min_select", "max_select"],
    "products": ["sku", "name", "category", "price", "is_available", "modifier_groups"],
    "channel_prices": ["channel_code", "sku", "price"],
    "outlet_prices": ["outlet_code", "sku", "price", "is_available"],
}
EXAMPLE = {
    "categories": ["Minuman", "bar", 1],
    "modifier_groups": ["Level Gula", "Normal", 0, 1, 1],
    "products": ["MN-1", "Es Teh", "Minuman", 5000, "ya", "Level Gula"],
    "channel_prices": ["GOFOOD", "MN-1", 6000],
    "outlet_prices": ["JKT01", "MN-1", "", "ya"],
}


@dataclass(frozen=True, slots=True)
class Counts:
    create: int = 0
    update: int = 0


@dataclass(slots=True)
class ImportResult:
    errors: list[RowError] = field(default_factory=list)
    categories: Counts = Counts()
    modifier_groups: Counts = Counts()
    products: Counts = Counts()
    channel_prices: int = 0
    outlet_prices: int = 0
    applied: bool = False

    def as_dict(self) -> dict[str, object]:
        return asdict(self)


def template() -> bytes:
    """One sheet per part of the menu, with the header row and one example row."""
    book = Workbook()
    book.remove(book.worksheets[0])
    for name in SHEETS:
        sheet = book.create_sheet(name)
        sheet.append(COLUMNS[name])
        sheet.append(EXAMPLE[name])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def read_workbook(data: bytes) -> dict[str, list[Row]]:
    """Rows as {column: value} by the header row; blank rows skipped; unknown sheets ignored."""
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, KeyError, ValueError, OSError) as err:
        raise DomainError("INVALID_WORKBOOK", "The file is not an .xlsx workbook") from err
    sheets: dict[str, list[Row]] = {}
    for name in SHEETS:
        if name not in book.sheetnames:
            continue
        rows = list(book[name].iter_rows(values_only=True))
        header = [
            str(h).strip().lower() if h is not None else "" for h in (rows[0] if rows else ())
        ]
        sheets[name] = [
            dict(zip(header, values, strict=False))
            for values in rows[1:]
            if any(v not in (None, "") for v in values)
        ]
    return sheets


async def _by_name[M: (CategoryModel, ModifierGroupModel)](
    uow: SqlAlchemyUnitOfWork, model: type[M], brand_id: str
) -> dict[str, M]:
    rows = await uow.session.scalars(
        select(model).where(model.brand_id == brand_id, model.deleted_at.is_(None))
    )
    return {r.name.lower(): r for r in rows}


async def import_menu(
    uow: SqlAlchemyUnitOfWork,
    new_id: Callable[[], str],
    sheets: dict[str, list[Row]],
    *,
    brand_id: str,
    outlets: dict[str, str],
    dry_run: bool,
) -> ImportResult:
    """`outlets` = the brand's outlets by upper-case code. Writes only when there is no error and
    `dry_run` is false; the caller then commits."""
    categories = await _by_name(uow, CategoryModel, brand_id)
    groups = await _by_name(uow, ModifierGroupModel, brand_id)
    products = {
        p.sku.lower(): p
        for p in await uow.session.scalars(
            select(ProductModel).where(
                ProductModel.brand_id == brand_id, ProductModel.deleted_at.is_(None)
            )
        )
    }
    brand_channels = {
        c.code.upper(): c for c in await channels.list_channels(uow, brand_id=brand_id)
    }
    parsed = parse_menu(
        sheets, known_categories=frozenset(categories), known_groups=frozenset(groups)
    )
    if isinstance(parsed, list):
        return ImportResult(errors=parsed)
    errors = [
        RowError("channel_prices", i + 2, f"Channel '{p.channel_code}' tidak ada")
        for i, p in enumerate(parsed.channel_prices)
        if p.channel_code not in brand_channels
    ] + [
        RowError("outlet_prices", i + 2, f"Outlet '{p.outlet_code}' tidak ada di brand ini")
        for i, p in enumerate(parsed.outlet_prices)
        if p.outlet_code not in outlets
    ]
    result = ImportResult(
        errors=errors,
        categories=_count([c.name.lower() for c in parsed.categories], categories),
        modifier_groups=_count([g.name.lower() for g in parsed.modifier_groups], groups),
        products=_count([p.sku.lower() for p in parsed.products], products),
        channel_prices=len(parsed.channel_prices),
        outlet_prices=len(parsed.outlet_prices),
    )
    if errors or dry_run:
        return result
    await _apply(
        uow, new_id, parsed, brand_id, outlets, {k: v.id for k, v in brand_channels.items()}
    )
    result.applied = True
    return result


def _count(keys: list[str], existing: Container[str]) -> Counts:
    updates = sum(1 for k in keys if k in existing)
    return Counts(create=len(keys) - updates, update=updates)


async def _apply_categories(
    uow: SqlAlchemyUnitOfWork, new_id: Callable[[], str], menu: MenuImport, brand_id: str
) -> dict[str, str]:
    existing = await _by_name(uow, CategoryModel, brand_id)
    for c in menu.categories:
        current = existing.get(c.name.lower())
        body = CategoryIn(
            name=c.name,
            print_role=c.print_role,
            sort_order=c.sort_order,
            color=current.color if current else "",
        )
        if current:
            await catalog.update_category(uow, current.id, body, brand_id=brand_id)
        else:
            await catalog.create_category(uow, new_id, body, brand_id=brand_id)
    return {k: v.id for k, v in (await _by_name(uow, CategoryModel, brand_id)).items()}


async def _apply_groups(
    uow: SqlAlchemyUnitOfWork, new_id: Callable[[], str], menu: MenuImport, brand_id: str
) -> dict[str, str]:
    existing = await _by_name(uow, ModifierGroupModel, brand_id)
    for g in menu.modifier_groups:
        current = existing.get(g.name.lower())
        kept = await _option_ids(uow, current.id) if current else {}
        body = ModifierGroupIn(
            name=g.name,
            min_select=g.min_select,
            max_select=g.max_select,
            options=[OptionIn(id=kept.get(n.lower()), name=n, price=p) for n, p in g.options],
        )
        if current:
            await catalog.update_modifier_group(uow, new_id, current.id, body, brand_id=brand_id)
        else:
            await catalog.create_modifier_group(uow, new_id, body, brand_id=brand_id)
    return {k: v.id for k, v in (await _by_name(uow, ModifierGroupModel, brand_id)).items()}


async def _apply_products(
    uow: SqlAlchemyUnitOfWork,
    new_id: Callable[[], str],
    menu: MenuImport,
    brand_id: str,
    categories: dict[str, str],
    groups: dict[str, str],
) -> dict[str, str]:
    existing = await _products_by_sku(uow, brand_id)
    ids: dict[str, str] = {}
    for p in menu.products:
        current = existing.get(p.sku.lower())
        body = ProductIn(
            category_id=categories[p.category.lower()],
            sku=p.sku,
            name=p.name,
            price=p.price,
            image_key=current.image_key if current else None,
            is_available=p.is_available,
            modifier_group_ids=[groups[n.lower()] for n in p.modifier_groups],
        )
        saved = (
            await catalog.update_product(uow, current.id, body, "", brand_id=brand_id)
            if current
            else await catalog.create_product(uow, new_id, body, "", brand_id=brand_id)
        )
        ids[p.sku.lower()] = saved.id
    return ids


async def _apply_channel_prices(
    uow: SqlAlchemyUnitOfWork,
    menu: MenuImport,
    brand_id: str,
    channel_ids: dict[str, str],
    product_ids: dict[str, str],
) -> None:
    """A channel's price list keeps the prices the file does not mention."""
    by_channel: dict[str, dict[str, int]] = {}
    for row in menu.channel_prices:
        by_channel.setdefault(row.channel_code, {})[product_ids[row.sku.lower()]] = row.price
    current = {c.id: c for c in await channels.list_channels(uow, brand_id=brand_id)}
    for code, prices in by_channel.items():
        channel_id = channel_ids[code]
        merged = {x.product_id: x.price for x in current[channel_id].prices} | prices
        body = ChannelPricesIn(
            prices=[ChannelPrice(product_id=k, price=v) for k, v in merged.items()]
        )
        await channels.set_channel_prices(uow, channel_id, body, brand_id=brand_id)


async def _apply(
    uow: SqlAlchemyUnitOfWork,
    new_id: Callable[[], str],
    menu: MenuImport,
    brand_id: str,
    outlets: dict[str, str],
    channel_ids: dict[str, str],
) -> None:
    categories = await _apply_categories(uow, new_id, menu, brand_id)
    groups = await _apply_groups(uow, new_id, menu, brand_id)
    product_ids = await _apply_products(uow, new_id, menu, brand_id, categories, groups)
    await _apply_channel_prices(uow, menu, brand_id, channel_ids, product_ids)
    for o in menu.outlet_prices:
        await catalog.set_outlet_product(
            uow,
            outlets[o.outlet_code],
            product_ids[o.sku.lower()],
            OutletProductIn(price=o.price, is_available=o.is_available),
            brand_id=brand_id,
        )


async def _products_by_sku(uow: SqlAlchemyUnitOfWork, brand_id: str) -> dict[str, ProductModel]:
    rows = await uow.session.scalars(
        select(ProductModel).where(
            ProductModel.brand_id == brand_id, ProductModel.deleted_at.is_(None)
        )
    )
    return {p.sku.lower(): p for p in rows}


async def _option_ids(uow: SqlAlchemyUnitOfWork, group_id: str) -> dict[str, str]:
    """Options keep their ids by name, so devices and past orders keep pointing at them."""
    rows = await uow.session.scalars(
        select(ModifierOptionModel).where(
            ModifierOptionModel.modifier_group_id == group_id,
            ModifierOptionModel.deleted_at.is_(None),
        )
    )
    return {r.name.lower(): r.id for r in rows}
