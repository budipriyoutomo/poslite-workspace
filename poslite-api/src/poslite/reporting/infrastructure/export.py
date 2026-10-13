"""Sales report downloads: CSV for any spreadsheet, XLSX with one sheet per breakdown."""

import csv
import io
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font

from poslite.reporting.infrastructure.queries import TOTAL_FIELDS, SalesSummary

DAY_HEADERS = [
    "Tanggal",
    "Penjualan kotor",
    "Diskon",
    "Service",
    "PB1",
    "Pendapatan",
    "Transaksi",
    "Void",
    "Refund",
    "Jumlah refund",
]
RUPIAH_FORMAT = "#,##0;-#,##0"


def _day_rows(summary: SalesSummary) -> list[list[object]]:
    rows: list[list[object]] = [
        [d["business_date"], *(d[f] for f in TOTAL_FIELDS)] for d in summary.days
    ]
    return [*rows, ["Total", *(summary.totals[f] for f in TOTAL_FIELDS)]]


def to_csv(summary: SalesSummary) -> bytes:
    """Daily rows + total; `;` separator and UTF-8 BOM so Indonesian-locale Excel opens it."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\n")
    writer.writerow(DAY_HEADERS)
    writer.writerows(_day_rows(summary))
    return buffer.getvalue().encode("utf-8-sig")


def _sheet(
    workbook: Workbook,
    title: str,
    headers: Sequence[str],
    rows: list[list[object]],
    money_columns: Sequence[int],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column in money_columns:
        for (cell,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            cell.number_format = RUPIAH_FORMAT
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = max(12, len(header) + 4)


def to_xlsx(summary: SalesSummary) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)  # type: ignore[arg-type]  # a new workbook always has one sheet
    _sheet(workbook, "Harian", DAY_HEADERS, _day_rows(summary), money_columns=[*range(2, 7), 9])
    _sheet(
        workbook,
        "Menu",
        ["Menu", "Jumlah", "Penjualan"],
        [[p["product_name"], p["quantity"], p["gross"]] for p in summary.products],
        [3],
    )
    _sheet(
        workbook,
        "Kategori",
        ["Kategori", "Jumlah", "Penjualan"],
        [[c["category_name"], c["quantity"], c["gross"]] for c in summary.categories],
        [3],
    )
    _sheet(
        workbook,
        "Pembayaran",
        ["Jenis", "Transaksi", "Jumlah"],
        [[p["method_type"], p["count"], p["amount"]] for p in summary.payment_methods],
        [3],
    )
    _sheet(
        workbook,
        "Kasir",
        ["Kasir", "Transaksi", "Pendapatan"],
        [[c["cashier_name"], c["order_count"], c["net"]] for c in summary.cashiers],
        [3],
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
