"""P1-4: the back-office downloads a template, previews a filled workbook with every problem per
row, then imports it in one transaction: all rows or none, matched by SKU and by name."""

import io
import time
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from ..identity.helpers import bearer
from ..sync.helpers import Setup, outlet_with_menu_and_device

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook(sheets: dict[str, list[list[Any]]]) -> bytes:
    book = Workbook()
    book.remove(book.active)  # type: ignore[arg-type]  # the default empty sheet
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


MENU = {
    "categories": [["name", "print_role", "sort_order"], ["Makanan", "kitchen", 1]],
    "modifier_groups": [
        ["group", "option", "price", "min_select", "max_select"],
        ["Topping", "Telur", 3000, 0, 2],
        ["Topping", "Keju", 4000, None, None],
    ],
    "products": [
        ["sku", "name", "category", "price", "is_available", "modifier_groups"],
        ["MK-1", "Mie Goreng", "Makanan", 20000, "ya", "Topping"],
        ["MN-1", "Es Teh Jumbo", "Minuman", 7000, "ya", None],  # the fixture's own SKU: updated
    ],
    "channel_prices": [["channel_code", "sku", "price"]],
    "outlet_prices": [
        ["outlet_code", "sku", "price", "is_available"],
        ["JKT01", "MK-1", 22000, "ya"],
    ],
}


async def _send(api: AsyncClient, s: Setup, body: bytes, *, dry_run: bool) -> Any:
    return await api.post(
        f"/api/v1/menu-import?dry_run={str(dry_run).lower()}",
        content=body,
        headers={**bearer(s.owner_token), "Content-Type": XLSX},
    )


async def _products(api: AsyncClient, s: Setup) -> dict[str, dict[str, Any]]:
    rows = (await api.get("/api/v1/products", headers=bearer(s.owner_token))).json()["data"]
    return {p["sku"]: p for p in rows}


async def test_the_template_has_a_sheet_per_part_of_the_menu(api: AsyncClient) -> None:
    s = await outlet_with_menu_and_device(api)

    response = await api.get("/api/v1/menu-import/template", headers=bearer(s.owner_token))

    assert response.headers["content-type"] == XLSX
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [
        "categories",
        "modifier_groups",
        "products",
        "channel_prices",
        "outlet_prices",
    ]
    assert [c.value for c in book["products"][1]] == [
        "sku",
        "name",
        "category",
        "price",
        "is_available",
        "modifier_groups",
    ]


async def test_a_preview_counts_what_would_change_and_writes_nothing(api: AsyncClient) -> None:
    s = await outlet_with_menu_and_device(api)

    preview = await _send(api, s, workbook(MENU), dry_run=True)

    assert preview.status_code == 200, preview.text
    assert preview.json()["data"] == {
        "errors": [],
        "categories": {"create": 1, "update": 0},
        "modifier_groups": {"create": 1, "update": 0},
        "products": {"create": 1, "update": 1},
        "channel_prices": 0,
        "outlet_prices": 1,
        "applied": False,
    }
    assert "MK-1" not in await _products(api, s)


async def test_imports_the_menu_and_a_second_import_updates_it(api: AsyncClient) -> None:
    s = await outlet_with_menu_and_device(api)

    first = await _send(api, s, workbook(MENU), dry_run=False)
    again = await _send(api, s, workbook(MENU), dry_run=False)

    assert first.json()["data"]["applied"] is True
    assert again.json()["data"]["products"] == {"create": 0, "update": 2}
    products = await _products(api, s)
    assert (products["MK-1"]["name"], products["MK-1"]["price"]) == ("Mie Goreng", 20000)
    assert (products["MN-1"]["name"], products["MN-1"]["price"]) == ("Es Teh Jumbo", 7000)
    assert len(products["MK-1"]["modifier_group_ids"]) == 1
    outlet = (
        await api.get(f"/api/v1/outlets/{s.outlet_id}/products", headers=bearer(s.owner_token))
    ).json()["data"]
    assert [p["effective_price"] for p in outlet if p["name"] == "Mie Goreng"] == [22000]


async def test_a_file_with_one_bad_row_imports_nothing(api: AsyncClient) -> None:
    s = await outlet_with_menu_and_device(api)
    bad = {**MENU, "products": [*MENU["products"], ["MK-2", "Nasi", "Makanan", -5, "ya", None]]}

    response = await _send(api, s, workbook(bad), dry_run=False)

    data = response.json()["data"]
    assert (response.status_code, data["applied"]) == (200, False)
    assert data["errors"] == [
        {"sheet": "products", "row": 4, "message": "Harga harus bilangan bulat ≥ 0"}
    ]
    assert "MK-1" not in await _products(api, s)


async def test_an_unknown_outlet_or_channel_is_a_row_error(api: AsyncClient) -> None:
    s = await outlet_with_menu_and_device(api)
    bad = {
        **MENU,
        "channel_prices": [["channel_code", "sku", "price"], ["GRAB", "MK-1", 1]],
        "outlet_prices": [
            ["outlet_code", "sku", "price", "is_available"],
            ["BDG9", "MK-1", 1, "ya"],
        ],
    }

    errors = (await _send(api, s, workbook(bad), dry_run=True)).json()["data"]["errors"]

    assert errors == [
        {"sheet": "channel_prices", "row": 2, "message": "Channel 'GRAB' tidak ada"},
        {"sheet": "outlet_prices", "row": 2, "message": "Outlet 'BDG9' tidak ada di brand ini"},
    ]


async def test_a_file_that_is_not_a_workbook_is_refused(api: AsyncClient) -> None:
    s = await outlet_with_menu_and_device(api)

    response = await _send(api, s, b"not a workbook", dry_run=True)

    assert (response.status_code, response.json()["code"]) == (422, "INVALID_WORKBOOK")


async def test_two_hundred_products_with_modifiers_import_quickly(api: AsyncClient) -> None:
    """DoD: a 200-product menu with modifiers goes in well under 5 minutes."""
    s = await outlet_with_menu_and_device(api)
    many = {
        **MENU,
        "products": [
            MENU["products"][0],
            *[[f"MK-{n}", f"Menu {n}", "Makanan", 10000 + n, "ya", "Topping"] for n in range(200)],
        ],
        "outlet_prices": [MENU["outlet_prices"][0]],
    }

    started = time.monotonic()
    response = await _send(api, s, workbook(many), dry_run=False)

    assert response.json()["data"]["products"] == {"create": 200, "update": 0}
    assert time.monotonic() - started < 30
