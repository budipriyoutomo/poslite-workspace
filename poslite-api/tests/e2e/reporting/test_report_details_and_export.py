import io
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine, async_sessionmaker

from poslite.reporting.api import rebuild
from poslite.worker import relay_job

from ..identity.helpers import bearer
from ..sync.helpers import Setup, outlet_with_menu_and_device
from ..sync.test_push import batch, push

RANGE = "from=2026-09-25&to=2026-09-25"


async def sold_and_projected(
    api: AsyncClient, app_engine: AsyncEngine, system_engine: AsyncEngine
) -> Setup:
    s = await outlet_with_menu_and_device(api)
    await push(api, s.device_token, batch(s))
    await relay_job(
        {
            "app_sessions": async_sessionmaker(app_engine),
            "system_sessions": async_sessionmaker(system_engine),
        }
    )
    return s


async def summary(api: AsyncClient, s: Setup) -> Any:
    return (
        await api.get(
            f"/api/v1/reports/sales-summary?outlet_id={s.outlet_id}&{RANGE}",
            headers=bearer(s.owner_token),
        )
    ).json()["data"]


async def test_summary_breaks_sales_down_by_category_and_cashier(
    api: AsyncClient, app_engine: AsyncEngine, system_engine: AsyncEngine
) -> None:
    s = await sold_and_projected(api, app_engine, system_engine)

    data = await summary(api, s)

    assert data["categories"] == [{"category_name": "Makanan", "quantity": 2, "gross": 70_000}]
    assert data["cashiers"] == [
        {"cashier_id": s.ids["siti"], "cashier_name": "Siti", "order_count": 1, "net": 72_765}
    ]


async def test_exports_csv_that_excel_opens_correctly(
    api: AsyncClient, app_engine: AsyncEngine, system_engine: AsyncEngine
) -> None:
    s = await sold_and_projected(api, app_engine, system_engine)

    response = await api.get(
        f"/api/v1/reports/sales-summary/export?outlet_id={s.outlet_id}&{RANGE}&format=csv",
        headers=bearer(s.owner_token),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert (
        'filename="penjualan-JKT01-2026-09-25-2026-09-25.csv"'
        in response.headers["content-disposition"]
    )
    assert response.content.startswith(b"\xef\xbb\xbf")  # BOM: Excel reads UTF-8
    assert response.content.decode("utf-8-sig").splitlines() == [
        "Tanggal;Penjualan kotor;Diskon;Service;PB1;Pendapatan;Transaksi;Void;Refund;Jumlah refund;"
        "Batal sebelum bayar;Nilai batal sebelum bayar;Pembulatan tunai",
        "2026-09-25;70000;7000;3150;6615;72765;1;0;0;0;0;0;0",
        "Total;70000;7000;3150;6615;72765;1;0;0;0;0;0;0",
    ]


async def test_exports_an_xlsx_workbook_with_a_sheet_per_breakdown(
    api: AsyncClient, app_engine: AsyncEngine, system_engine: AsyncEngine
) -> None:
    s = await sold_and_projected(api, app_engine, system_engine)

    response = await api.get(
        f"/api/v1/reports/sales-summary/export?outlet_id={s.outlet_id}&{RANGE}&format=xlsx",
        headers=bearer(s.owner_token),
    )

    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Harian", "Menu", "Kategori", "Pembayaran", "Kasir"]
    assert [c.value for c in workbook["Harian"][2]] == [
        "2026-09-25",
        70_000,
        7_000,
        3_150,
        6_615,
        72_765,
        1,
        0,
        0,
        0,
        0,
        0,
        0,
    ]
    assert workbook["Menu"]["A2"].value == "Nasi Gudeg"


async def test_rebuild_recomputes_the_read_model_from_the_orders(
    api: AsyncClient, app_engine: AsyncEngine, system_engine: AsyncEngine
) -> None:
    s = await sold_and_projected(api, app_engine, system_engine)
    async with system_engine.begin() as conn:
        for table in (
            "rpt_daily_sales",
            "rpt_daily_product_sales",
            "rpt_daily_payment_sales",
            "rpt_daily_category_sales",
            "rpt_daily_cashier_sales",
        ):
            await conn.execute(text(f"DELETE FROM {table}"))  # noqa: S608 — fixed table names

    rebuilt = await rebuild(
        system=async_sessionmaker(system_engine),
        app=async_sessionmaker(app_engine),
        start="2026-09-25",
        end="2026-09-25",
    )

    assert rebuilt == 1  # one outlet-day
    data = await summary(api, s)
    assert (data["totals"]["net"], data["cashiers"][0]["net"]) == (72_765, 72_765)
